"""
Genera la matriz de Google Sheets de Sarah & Tin (archivo .xlsx para subir a
Google Drive) a partir de lib/sync/sheet-schema.json, la misma especificación
que usa el motor de sincronización.

Uso:  python scripts/build-matrix.py [salida.xlsx]

La planilla sale VACÍA de datos: los registros llegan desde la app con la
primera sincronización. No contiene datos reales ni credenciales.
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
SPEC = json.loads((ROOT / "lib" / "sync" / "sheet-schema.json").read_text(encoding="utf8"))
OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "docs" / "Matriz_Sarah_Tin.xlsx"

MAX_ROWS = 2000  # filas cubiertas por validaciones y formato condicional
FORMATS = {
    "money": '"$"#,##0',
    "pct": "0.0%",
    "int": "#,##0",
    "num": "#,##0.###",
    "date": "DD-MM-YYYY",
    "datetime": "DD-MM-YYYY HH:MM",
    "time": "HH:MM",
    "month": "MMM YYYY",
}
FORMAT_ROWS = 300  # filas con formato de número/fecha ya aplicado (Preparar matriz extiende el formato)

WHITE = "FFFFFF"
COCOA = "5B3A2E"
READONLY_FILL = PatternFill("solid", fgColor="EFEAE8")
TECH_FILL = PatternFill("solid", fgColor="D9D4D2")
BOLD_WHITE = Font(bold=True, color=WHITE)
BOLD_DARK = Font(bold=True, color=COCOA)
TECH_FONT = Font(bold=True, color="6B6461", italic=True)

TABS = {t["name"]: t for t in SPEC["tabs"]}


def col_of(tab, key):
    keys = [c["key"] for c in TABS[tab]["columns"]]
    return get_column_letter(keys.index(key) + 1)


def resolve(formula):
    """{Tab.key} → Tab!X2:X  ·  {Tab.a:b} → Tab!X2:Y"""
    def rep(m):
        tab, keys = m.group(1), m.group(2)
        if ":" in keys:
            a, b = keys.split(":")
            return f"{tab}!{col_of(tab, a)}2:{col_of(tab, b)}"
        c = col_of(tab, keys)
        return f"{tab}!{c}2:{c}"
    return re.sub(r"\{([A-Za-z_]+)\.([a-z_:]+)\}", rep, formula)


def group_fill(group):
    return PatternFill("solid", fgColor=SPEC["groups"][group]["color"])


def header_row(ws, headers, fills, fonts):
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = fills[i - 1]
        c.font = fonts[i - 1]
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"


def build_sync_tab(wb, tab):
    ws = wb.create_sheet(tab["name"])
    ws.sheet_properties.tabColor = SPEC["groups"][tab["group"]]["color"]
    cols = tab["columns"]
    tech = SPEC["technical_columns"]
    headers = [c["header"] for c in cols] + [t["header"] for t in tech]
    fills, fonts = [], []
    for c in cols:
        if c.get("edit"):
            fills.append(group_fill(tab["group"]))
            fonts.append(BOLD_WHITE)
        else:
            fills.append(READONLY_FILL)
            fonts.append(BOLD_DARK)
    fills += [TECH_FILL] * len(tech)
    fonts += [TECH_FONT] * len(tech)
    header_row(ws, headers, fills, fonts)

    for i, c in enumerate(cols, start=1):
        letter = get_column_letter(i)
        width = c.get("width") or (38 if c["type"] == "id" else max(12, min(30, len(c["header"]) + 4)))
        ws.column_dimensions[letter].width = width
        rng = f"{letter}2:{letter}{MAX_ROWS}"
        if not c.get("edit"):
            ws.conditional_formatting.add(rng, FormulaRule(formula=["TRUE"], fill=READONLY_FILL))
        fmt = FORMATS.get(c["type"])
        if fmt:
            for r in range(2, FORMAT_ROWS + 2):
                ws.cell(row=r, column=i).number_format = fmt
        elif c["type"] in ("text", "code", "month", "ref", "id"):
            for r in range(2, FORMAT_ROWS + 2):
                ws.cell(row=r, column=i).number_format = "@"
        if c.get("edit") and c["type"] == "enum":
            dv = DataValidation(type="list", formula1='"' + ",".join(c["options"]) + '"', allow_blank=not c.get("required"),
                                showErrorMessage=True, errorTitle="Valor no válido",
                                error="Elige una opción de la lista: " + ", ".join(c["options"]))
            ws.add_data_validation(dv)
            dv.add(rng)
        if c.get("edit") and c["type"] == "bool":
            dv = DataValidation(type="list", formula1='"TRUE,FALSE"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(rng)
        if c.get("edit") and c["type"] == "ref":
            ref_col = col_of(c["ref"], "id")
            dv = DataValidation(type="list", formula1=f"={c['ref']}!${ref_col}$2:${ref_col}${MAX_ROWS}", allow_blank=True,
                                showErrorMessage=True, errorTitle="ID no encontrado",
                                error=f"Copia el ID desde la hoja {c['ref']}.")
            ws.add_data_validation(dv)
            dv.add(rng)
    for j, t in enumerate(tech, start=len(cols) + 1):
        letter = get_column_letter(j)
        ws.column_dimensions[letter].width = 14 if t["key"] != "_estado_sync" else 26
        ws.conditional_formatting.add(f"{letter}2:{letter}{MAX_ROWS}", FormulaRule(formula=["TRUE"], fill=TECH_FILL))

    # Estado de sincronización resaltado.
    st = get_column_letter(len(cols) + 3)
    ws.conditional_formatting.add(f"{st}2:{st}{MAX_ROWS}", FormulaRule(
        formula=[f'LEFT({st}2,5)="ERROR"'], fill=PatternFill("solid", fgColor="F6C9C9"), font=Font(color="9B1C1C", bold=True)))
    ws.conditional_formatting.add(f"{st}2:{st}{MAX_ROWS}", FormulaRule(
        formula=[f'{st}2="CONFLICTO"'], fill=PatternFill("solid", fgColor="FBE3B5"), font=Font(color="7A4B00", bold=True)))

    if tab.get("kind") == "params":
        for r, p in enumerate(tab["params"], start=2):
            ws.cell(row=r, column=1, value=p["key"])
            ws.cell(row=r, column=2, value=p["label"])
            ws.cell(row=r, column=4, value=p.get("help", ""))
    return ws


def build_analysis_tab(wb, a):
    ws = wb.create_sheet(a["name"])
    ws.sheet_properties.tabColor = SPEC["groups"]["C"]["color"]
    n = len(a["headers"])
    header_row(ws, a["headers"], [group_fill("C")] * n, [BOLD_WHITE] * n)
    ws.freeze_panes = "A2"
    for i, fmt in enumerate(a["formats"], start=1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = 22 if i == 1 else 16
        f = FORMATS.get(fmt)
        if f:
            for r in range(2, 202):
                ws.cell(row=r, column=i).number_format = f
    if "row_formulas" in a:
        for r in range(2, a["rows"] + 2):
            for i, f in enumerate(a["row_formulas"], start=1):
                ws.cell(row=r, column=i, value=resolve(f.replace("{r}", str(r))))
        tr = a["rows"] + 2
        for i, f in enumerate(a.get("total_row", []), start=1):
            c = ws.cell(row=tr, column=i, value=resolve(f) if f.startswith("=") else f)
            c.font = Font(bold=True)
    for i, f in enumerate(a.get("array_formulas", []), start=1):
        if f:
            ws.cell(row=2, column=i, value=resolve(f))
    # Nota de ayuda a la derecha de la tabla.
    note = ws.cell(row=1, column=n + 2, value=a["help"])
    note.alignment = Alignment(wrap_text=True, vertical="top")
    note.font = Font(italic=True, color="6B6461")
    ws.column_dimensions[get_column_letter(n + 2)].width = 60
    return ws


def build_readme(wb):
    ws = wb.active
    ws.title = "Leer_Primero"
    ws.sheet_properties.tabColor = COCOA
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 110
    lines = [
        ("t", SPEC["title"]),
        ("p", "Planilla maestra de configuración y análisis. La app (celular de Camila) registra la operación diaria; esta planilla se mantiene sincronizada con la app."),
        ("h", "Cómo se organiza"),
        ("p", "Pestañas ROSADAS (Grupo A · Configuración): César edita aquí precios, costos, recetas, parámetros, listas. Se aplican a la app al sincronizar."),
        ("p", "Pestañas AZULES (Grupo B · Operación): pedidos, ventas, cobros, gastos, compras, producción. Se llenan desde la app; aquí solo se corrigen notas y algunos textos."),
        ("p", "Pestañas DORADAS (Grupo C · Análisis): fórmulas automáticas (resumen mensual, rentabilidad, por cobrar, por pagar, alertas). No escribir en ellas."),
        ("h", "Colores de columnas"),
        ("p", "Encabezado de color = columna editable.  Encabezado y celdas grises = calculado por la app (si lo cambias, se restaura).  Columnas que empiezan con _ = técnicas, no tocar."),
        ("h", "Reglas para no perder datos"),
        ("p", "1. No cambies ni muevas los encabezados (fila 1). Si cambian, esa pestaña deja de sincronizar y aparece un error en la app."),
        ("p", "2. No borres filas: para dejar de usar algo, desmarca la columna Activo. Una fila borrada vuelve a aparecer en la siguiente sincronización."),
        ("p", "3. Para agregar un ingrediente, producto, receta, cliente u opción: escribe una fila nueva al final y deja el ID vacío. La app le asigna ID."),
        ("p", "4. No copies filas con su ID: un ID duplicado queda marcado como ERROR."),
        ("p", "5. Fechas en formato DD-MM-AAAA. Montos en pesos, sin decimales. Porcentajes como 60%."),
        ("p", "6. Para relacionar (por ejemplo una línea de receta con su ingrediente) copia el ID desde la pestaña correspondiente."),
        ("h", "Columna _estado_sync"),
        ("p", "OK = sincronizado.  PENDIENTE = alta en proceso.  ERROR: motivo = no se aplicó; corrige el valor y se reintenta solo.  CONFLICTO = la fila cambió aquí y en la app al mismo tiempo: César elige qué versión queda en la app (Configuración → Sincronización)."),
        ("h", "Cuándo se sincroniza"),
        ("p", "Cada 10 minutos, después de cada registro en la app y al presionar \"Sincronizar ahora\" en la app (Configuración → Sincronización). La pestaña Sync_Bitacora muestra los últimos movimientos."),
        ("h", "Conceptos financieros"),
        ("p", "Venta = lo vendido (aunque se cobre después).  Cobro/abono = dinero recibido.  Reembolso = devolución (monto negativo).  Compra de insumos = entra a inventario; no es gasto.  Gasto = costo operativo (luz, gas, transporte…).  Costo de lo vendido = costo histórico guardado al vender.  Aporte / retiro = dinero de las dueñas, no es venta ni gasto.  Herramienta de gestión interna: no es contabilidad tributaria."),
        ("h", "Diccionario"),
        ("p", "La pestaña Diccionario describe cada columna de cada pestaña."),
    ]
    r = 2
    for kind, text in lines:
        c = ws.cell(row=r, column=2, value=text)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "t":
            c.font = Font(bold=True, size=18, color="D24D77")
        elif kind == "h":
            c.font = Font(bold=True, size=12, color=COCOA)
            r += 0
        r += 1 if kind != "h" else 1
        if kind == "h":
            ws.row_dimensions[r - 1].height = 24
    return ws


def build_dictionary(wb):
    ws = wb.create_sheet("Diccionario")
    ws.sheet_properties.tabColor = COCOA
    headers = ["Pestaña", "Grupo", "Columna", "Tipo", "Editable en Sheets", "Obligatoria", "Explicación"]
    header_row(ws, headers, [PatternFill("solid", fgColor=COCOA)] * 7, [BOLD_WHITE] * 7)
    types = {"id": "ID", "text": "Texto", "code": "Código", "int": "Entero", "num": "Número", "money": "Monto $",
             "pct": "Porcentaje", "bool": "Sí/No", "date": "Fecha", "datetime": "Fecha y hora", "time": "Hora",
             "month": "Mes AAAA-MM", "enum": "Lista", "ref": "ID de otra pestaña"}
    r = 2
    for t in SPEC["tabs"]:
        for c in t["columns"]:
            help_ = c.get("help", "")
            if c["type"] == "enum":
                help_ = (help_ + " Opciones: " + ", ".join(c["options"])).strip()
            if c["type"] == "ref":
                help_ = (help_ + f" ID de la pestaña {c['ref']}.").strip()
            ws.append([t["name"], SPEC["groups"][t["group"]]["name"], c["header"], types[c["type"]],
                       "Sí" if c.get("edit") else "No", "Sí" if c.get("required") else "", help_])
            r += 1
    for t in SPEC["technical_columns"]:
        ws.append(["(todas)", "", t["header"], "Técnica", "No", "", t["help"]])
    for i, w in enumerate([18, 14, 28, 16, 12, 11, 80], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    return ws


def main():
    wb = Workbook()
    build_readme(wb)
    for t in SPEC["tabs"]:
        if t["group"] != "C":
            build_sync_tab(wb, t)
    for a in SPEC["analysis"]:
        build_analysis_tab(wb, a)
    for t in SPEC["tabs"]:
        if t["group"] == "C":
            build_sync_tab(wb, t)
    build_dictionary(wb)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"{OUT}  ({OUT.stat().st_size // 1024} KB, {len(wb.sheetnames)} pestañas)")


if __name__ == "__main__":
    main()
